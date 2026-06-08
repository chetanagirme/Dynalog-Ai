import chromadb
import os
import re
import PyPDF2
from docx import Document

chroma_client = chromadb.PersistentClient(path="data/chroma")

def get_collection(project_id):
    return chroma_client.get_or_create_collection(
        name=f"project_{project_id}",
        metadata={"hnsw:space": "cosine"}
    )

def read_file(filepath):
    ext = filepath.rsplit('.', 1)[-1].lower()
    try:
        if ext in ['txt', 'csv', 'md', 'log', 'json', 'xml', 'html', 'htm', 'py', 'js', 'c', 'cpp', 'h']:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        elif ext == 'pdf':
            text = ''
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for i, page in enumerate(reader.pages):
                    t = page.extract_text()
                    if t:
                        text += f"\n[Page {i+1}]\n{t}"
            return text
        elif ext in ['docx', 'doc']:
            doc = Document(filepath)
            return '\n'.join([p.text for p in doc.paragraphs if p.text.strip()])
        elif ext in ['xlsx', 'xls']:
            try:
                import openpyxl
                wb   = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
                text = ''
                for sheet in wb.worksheets:
                    text += f"\n[Sheet: {sheet.title}]\n"
                    for row in sheet.iter_rows(values_only=True):
                        row_text = ' | '.join([str(c) for c in row if c is not None])
                        if row_text.strip():
                            text += row_text + '\n'
                return text
            except:
                with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
        elif ext == 'pptx':
            try:
                from pptx import Presentation
                prs  = Presentation(filepath)
                text = ''
                for i, slide in enumerate(prs.slides):
                    text += f"\n[Slide {i+1}]\n"
                    for shape in slide.shapes:
                        if hasattr(shape, 'text') and shape.text.strip():
                            text += shape.text + '\n'
                return text
            except:
                return f"Could not read PPTX: {os.path.basename(filepath)}"
        else:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()
                if len(content.strip()) > 10:
                    return content
                return f"Binary file: {os.path.basename(filepath)}"
    except Exception as e:
        return f"Could not read file: {str(e)}"
    return ""

def chunk_text(text, filename, chunk_size=500, overlap=50):
    sentences = re.split(r'(?<=[.!?])\s+', text)
    chunks    = []
    current   = ""
    for sentence in sentences:
        if len(current) + len(sentence) < chunk_size:
            current += " " + sentence
        else:
            if current.strip():
                chunks.append(current.strip())
            current = sentence
    if current.strip():
        chunks.append(current.strip())
    return [c for c in chunks if len(c) > 30]

def index_file(project_id, filepath, filename):
    try:
        collection = get_collection(project_id)
        text       = read_file(filepath)
        if not text or len(text.strip()) < 10:
            return False, "File is empty or unreadable"
        chunks = chunk_text(text, filename)
        if not chunks:
            return False, "No content to index"
        ids       = [f"{project_id}_{filename}_{i}" for i in range(len(chunks))]
        metadatas = [{"filename": filename, "chunk": i, "total": len(chunks)} for i in range(len(chunks))]
        try:
            existing = collection.get(ids=[ids[0]])
            if existing['ids']:
                all_ids    = collection.get()['ids']
                to_delete  = [id for id in all_ids if id.startswith(f"{project_id}_{filename}_")]
                if to_delete:
                    collection.delete(ids=to_delete)
        except:
            pass
        collection.add(documents=chunks, metadatas=metadatas, ids=ids)
        return True, len(chunks)
    except Exception as e:
        return False, str(e)

def search_project(project_id, query, n_results=5):
    try:
        collection = get_collection(project_id)
        count      = collection.count()
        if count == 0:
            return []
        results = collection.query(
            query_texts=[query],
            n_results=min(n_results, count)
        )
        hits = []
        if results['documents'] and results['documents'][0]:
            for doc, meta, dist in zip(
                results['documents'][0],
                results['metadatas'][0],
                results['distances'][0]
            ):
                hits.append({
                    'content':  doc,
                    'filename': meta.get('filename', 'Unknown'),
                    'chunk':    meta.get('chunk', 0),
                    'score':    round(1 - dist, 3)
                })
        return hits
    except Exception as e:
        return []

def delete_project_index(project_id):
    try:
        chroma_client.delete_collection(f"project_{project_id}")
    except:
        pass

def delete_file_from_index(project_id, filename):
    try:
        collection = get_collection(project_id)
        all_ids    = collection.get()['ids']
        to_delete  = [id for id in all_ids if id.startswith(f"{project_id}_{filename}_")]
        if to_delete:
            collection.delete(ids=to_delete)
    except:
        pass