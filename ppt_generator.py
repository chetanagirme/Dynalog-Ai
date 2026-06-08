from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import ollama
import json
import io
import re
import datetime

DARK_BLUE  = RGBColor(10,  22,  40)
MID_BLUE   = RGBColor(26,  46,  74)
ACCENT     = RGBColor(55, 138, 221)
WHITE      = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(245, 245, 243)
MID_GRAY   = RGBColor(100, 100, 100)

def add_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, l, t, w, h, color):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_text(slide, text, l, t, w, h, size, color, bold=False, align=PP_ALIGN.LEFT, italic=False):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf  = txb.text_frame
    tf.word_wrap = True
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.color.rgb = color
    run.font.bold  = bold
    run.font.italic = italic
    return txb

def add_bullet_slide(prs, title, bullets, slide_num, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, DARK_BLUE)
    add_rect(slide, 0, 0, 10, 0.08, ACCENT)
    add_rect(slide, 0, 6.92, 10, 0.08, ACCENT)
    add_text(slide, title, 0.4, 0.15, 9.2, 0.7, 24, WHITE, bold=True)
    add_rect(slide, 0.4, 0.95, 9.2, 0.03, ACCENT)
    y = 1.1
    for bullet in bullets[:7]:
        if not bullet.strip():
            continue
        dot = slide.shapes.add_shape(1, Inches(0.4), Inches(y+0.08), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACCENT
        dot.line.fill.background()
        add_text(slide, bullet.strip().lstrip('•-– '), 0.65, y, 8.9, 0.45, 14, LIGHT_GRAY)
        y += 0.52
        if y > 6.5:
            break
    add_text(slide, f'{slide_num} / {total}', 9.0, 7.0, 0.8, 0.3, 10, MID_GRAY, align=PP_ALIGN.RIGHT)

def add_two_col_slide(prs, title, left_items, right_items, slide_num, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, DARK_BLUE)
    add_rect(slide, 0, 0, 10, 0.08, ACCENT)
    add_rect(slide, 0, 6.92, 10, 0.08, ACCENT)
    add_text(slide, title, 0.4, 0.15, 9.2, 0.7, 24, WHITE, bold=True)
    add_rect(slide, 0.4, 0.95, 9.2, 0.03, ACCENT)
    add_rect(slide, 0.4,  1.1, 4.4, 5.6, MID_BLUE)
    add_rect(slide, 5.2,  1.1, 4.4, 5.6, MID_BLUE)
    y = 1.3
    for item in left_items[:6]:
        if item.strip():
            add_text(slide, '• ' + item.strip().lstrip('•-– '), 0.55, y, 4.1, 0.45, 13, LIGHT_GRAY)
            y += 0.5
    y = 1.3
    for item in right_items[:6]:
        if item.strip():
            add_text(slide, '• ' + item.strip().lstrip('•-– '), 5.35, y, 4.1, 0.45, 13, LIGHT_GRAY)
            y += 0.5
    add_text(slide, f'{slide_num} / {total}', 9.0, 7.0, 0.8, 0.3, 10, MID_GRAY, align=PP_ALIGN.RIGHT)

def get_model():
    try:
        model_list = ollama.list().get('models', [])
        available = []
        for m in model_list:
            try:
                name = m.get('name') or m.get('model') or ''
                available.append(name.split(':')[0])
            except:
                continue
        print(f"[PPT] Available models: {available}")
        for preferred in ['llama3', 'mistral', 'phi3', 'llama2']:
            if preferred in available:
                return preferred
    except Exception as e:
        print(f"[PPT] Model list error: {e}")
    return 'llama3'

def generate_ppt(project_name, project_desc, search_results, output_path=None):
    context = '\n\n'.join([
        f"[{h['filename']}]: {h['content']}"
        for h in search_results[:10]
    ])

    today = datetime.date.today().strftime('%B %Y')

    prompt = f"""Create a PowerPoint presentation for project: "{project_name}".
Description: {project_desc}

Content from project files:
{context[:3000]}

Return ONLY this JSON, nothing else, no markdown backticks:
{{
  "title": "{project_name}",
  "subtitle": "Technical Project Overview",
  "date": "{today}",
  "slides": [
    {{
      "type": "bullets",
      "title": "Project Overview",
      "bullets": ["point 1", "point 2", "point 3", "point 4"]
    }},
    {{
      "type": "bullets",
      "title": "Key Specifications",
      "bullets": ["spec 1", "spec 2", "spec 3", "spec 4"]
    }},
    {{
      "type": "bullets",
      "title": "System Architecture",
      "bullets": ["component 1", "component 2", "component 3", "component 4"]
    }},
    {{
      "type": "two_col",
      "title": "Technical Details",
      "left": ["detail 1", "detail 2", "detail 3"],
      "right": ["detail 4", "detail 5", "detail 6"]
    }},
    {{
      "type": "bullets",
      "title": "Key Findings",
      "bullets": ["finding 1", "finding 2", "finding 3"]
    }},
    {{
      "type": "bullets",
      "title": "Conclusion",
      "bullets": ["conclusion 1", "conclusion 2", "conclusion 3"]
    }}
  ]
}}

Replace ALL placeholder text with REAL content extracted from the project files above.
Return ONLY the JSON object. No explanation. No markdown."""

    model = get_model()
    print(f"[PPT] Using model: {model}")

    reply = None
    try:
        response = ollama.chat(
            model=model,
            messages=[{'role': 'user', 'content': prompt}],
            options={'num_predict': 1200, 'temperature': 0.1}
        )
        reply = response['message']['content']
        print(f"[PPT] Got reply: {reply[:150]}")
    except Exception as e:
        print(f"[PPT] Model error: {e}")
        raise Exception(f"AI model failed: {str(e)}. Make sure Ollama is running.")

    # Extract JSON from reply
    json_match = re.search(r'\{[\s\S]*\}', reply)
    if not json_match:
        print(f"[PPT] No JSON found in reply: {reply}")
        raise Exception("AI did not return valid JSON. Try again.")

    try:
        data = json.loads(json_match.group())
    except json.JSONDecodeError as e:
        print(f"[PPT] JSON parse error: {e}")
        # Use fallback data
        data = {
            "title": project_name,
            "subtitle": "Project Overview",
            "date": today,
            "slides": [
                {"type": "bullets", "title": "Project Overview",
                 "bullets": ["Project: " + project_name, project_desc or "No description", "Generated by Dynalog AI", ""]},
                {"type": "bullets", "title": "Document Content",
                 "bullets": [h['content'][:80] for h in search_results[:4]]},
                {"type": "bullets", "title": "Key Information",
                 "bullets": [h['filename'] for h in search_results[:4]]},
                {"type": "bullets", "title": "Conclusion",
                 "bullets": ["Review complete", "See attached documents", "Prepared by Dynalog AI"]}
            ]
        }

    # Build presentation
    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(7.5)

    total_slides = len(data.get('slides', [])) + 2

    # Title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(title_slide, DARK_BLUE)
    add_rect(title_slide, 0, 0,   10, 0.08, ACCENT)
    add_rect(title_slide, 0, 2.8, 10, 2.0,  MID_BLUE)
    add_rect(title_slide, 0, 7.42,10, 0.08, ACCENT)
    add_text(title_slide, 'DYNALOG DEFENCE ORGANISATION',
             0.5, 0.3, 9, 0.5, 11, ACCENT, bold=True, align=PP_ALIGN.CENTER)
    add_text(title_slide, data.get('title', project_name),
             0.5, 2.95, 9, 0.9, 28, WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(title_slide, data.get('subtitle', ''),
             0.5, 3.9, 9, 0.5, 16, LIGHT_GRAY, align=PP_ALIGN.CENTER)
    add_rect(title_slide, 4.5, 4.55, 1.0, 0.04, ACCENT)
    add_text(title_slide, data.get('date', today),
             0.5, 4.7, 9, 0.4, 13, MID_GRAY, align=PP_ALIGN.CENTER)
    add_text(title_slide, 'CONFIDENTIAL - INTERNAL USE ONLY',
             0.5, 7.1, 9, 0.3, 10, MID_GRAY, align=PP_ALIGN.CENTER, italic=True)

    # Content slides
    slide_num = 2
    for slide_data in data.get('slides', []):
        stype = slide_data.get('type', 'bullets')
        if stype == 'two_col':
            add_two_col_slide(
                prs,
                slide_data.get('title', 'Details'),
                slide_data.get('left', []),
                slide_data.get('right', []),
                slide_num, total_slides
            )
        else:
            add_bullet_slide(
                prs,
                slide_data.get('title', 'Slide'),
                slide_data.get('bullets', []),
                slide_num, total_slides
            )
        slide_num += 1

    # Thank you slide
    end_slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(end_slide, DARK_BLUE)
    add_rect(end_slide, 0, 0,    10, 0.08, ACCENT)
    add_rect(end_slide, 0, 7.42, 10, 0.08, ACCENT)
    add_rect(end_slide, 0, 2.9,  10, 1.7,  MID_BLUE)
    add_text(end_slide, 'Thank You',
             0.5, 3.0, 9, 0.8, 40, WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(end_slide, project_name,
             0.5, 3.85, 9, 0.5, 16, ACCENT, align=PP_ALIGN.CENTER)
    add_text(end_slide, 'CONFIDENTIAL - INTERNAL USE ONLY',
             0.5, 7.1, 9, 0.3, 10, MID_GRAY, align=PP_ALIGN.CENTER, italic=True)

    if output_path:
        prs.save(output_path)
        return output_path
    else:
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        return buf