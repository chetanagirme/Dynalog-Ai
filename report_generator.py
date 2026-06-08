import ollama
import json
import re
import io
import datetime

def get_model():
    try:
        model_list = ollama.list().get('models', [])
        available = []
        for m in model_list:
            try:
                name = m.get('name') or m.get('model') or ''
                available.append(name.split(':')[0])
            except:
                continue
        for preferred in ['llama3', 'mistral', 'phi3', 'llama2']:
            if preferred in available:
                return preferred
    except:
        pass
    return 'llama3'

def ask_ai(prompt, max_tokens=1500):
    model = get_model()
    try:
        response = ollama.chat(
            model=model,
            messages=[{'role': 'user', 'content': prompt}],
            options={'num_predict': max_tokens, 'temperature': 0.1}
        )
        return response['message']['content']
    except Exception as e:
        raise Exception(f"AI error: {str(e)}")

def generate_overview(project_name, search_results):
    context = '\n\n'.join([f"[{h['filename']}]: {h['content']}" for h in search_results[:8]])
    prompt  = f"""Analyse this project: "{project_name}"

Content:
{context[:3000]}

Return ONLY JSON, no markdown:
{{
  "abstract": "2-3 sentence project summary",
  "key_topics": ["topic1", "topic2", "topic3", "topic4", "topic5"],
  "main_findings": ["finding1", "finding2", "finding3"],
  "important_dates": ["date/event 1", "date/event 2"],
  "project_type": "Electronics/Defence/Software/etc",
  "complexity": "Low/Medium/High",
  "status": "Active/Completed/In Progress"
}}"""
    reply = ask_ai(prompt, 600)
    match = re.search(r'\{[\s\S]*\}', reply)
    if match:
        try:
            return json.loads(match.group())
        except:
            pass
    return {
        "abstract": f"This project '{project_name}' contains technical documentation and related files.",
        "key_topics": ["System Design", "Technical Specs", "Testing", "Components", "Architecture"],
        "main_findings": ["Documentation indexed successfully", "Files available for analysis"],
        "important_dates": [],
        "project_type": "Electronics",
        "complexity": "Medium",
        "status": "Active"
    }

def generate_faq(project_name, search_results):
    context = '\n\n'.join([f"[{h['filename']}]: {h['content']}" for h in search_results[:8]])
    prompt  = f"""Project: "{project_name}"

Content:
{context[:3000]}

Generate 8 frequently asked questions with answers based on the content above.
Return ONLY JSON, no markdown:
{{
  "faqs": [
    {{"q": "question here?", "a": "answer here"}},
    {{"q": "question here?", "a": "answer here"}}
  ]
}}"""
    reply = ask_ai(prompt, 1000)
    match = re.search(r'\{[\s\S]*\}', reply)
    if match:
        try:
            return json.loads(match.group())
        except:
            pass
    return {"faqs": [{"q": "What is this project about?", "a": f"This is the {project_name} project containing technical documentation."}]}

def generate_flowchart(project_name, search_results):
    context = '\n\n'.join([f"[{h['filename']}]: {h['content']}" for h in search_results[:6]])
    prompt  = f"""Project: "{project_name}"

Content:
{context[:2500]}

Generate a process flowchart for this project.
Return ONLY JSON, no markdown:
{{
  "title": "Process flowchart title",
  "nodes": [
    {{"id": "1", "label": "Start", "type": "start"}},
    {{"id": "2", "label": "Process step", "type": "process"}},
    {{"id": "3", "label": "Decision?", "type": "decision"}},
    {{"id": "4", "label": "End", "type": "end"}}
  ],
  "edges": [
    {{"from": "1", "to": "2", "label": ""}},
    {{"from": "2", "to": "3", "label": ""}},
    {{"from": "3", "to": "4", "label": "Yes"}}
  ]
}}

Types: start, end, process, decision, io
Generate 6-10 nodes based on actual project content. Keep labels SHORT (max 4 words)."""
    reply = ask_ai(prompt, 800)
    match = re.search(r'\{[\s\S]*\}', reply)
    if match:
        try:
            return json.loads(match.group())
        except:
            pass
    return {
        "title": f"{project_name} Process Flow",
        "nodes": [
            {"id":"1","label":"Start","type":"start"},
            {"id":"2","label":"Input Data","type":"io"},
            {"id":"3","label":"Process","type":"process"},
            {"id":"4","label":"Valid?","type":"decision"},
            {"id":"5","label":"Output","type":"io"},
            {"id":"6","label":"End","type":"end"}
        ],
        "edges": [
            {"from":"1","to":"2","label":""},
            {"from":"2","to":"3","label":""},
            {"from":"3","to":"4","label":""},
            {"from":"4","to":"5","label":"Yes"},
            {"from":"5","to":"6","label":""}
        ]
    }

def generate_mindmap(project_name, search_results):
    context = '\n\n'.join([f"[{h['filename']}]: {h['content']}" for h in search_results[:6]])
    prompt  = f"""Project: "{project_name}"

Content:
{context[:2500]}

Generate a mind map showing how concepts relate.
Return ONLY JSON, no markdown:
{{
  "center": "{project_name}",
  "branches": [
    {{
      "topic": "Main Topic 1",
      "color": "blue",
      "subtopics": ["subtopic1", "subtopic2", "subtopic3"]
    }},
    {{
      "topic": "Main Topic 2",
      "color": "green",
      "subtopics": ["subtopic1", "subtopic2"]
    }}
  ]
}}

Colors available: blue, green, amber, purple, coral
Generate 5-6 main branches with 2-4 subtopics each based on actual project content."""
    reply = ask_ai(prompt, 800)
    match = re.search(r'\{[\s\S]*\}', reply)
    if match:
        try:
            return json.loads(match.group())
        except:
            pass
    return {
        "center": project_name,
        "branches": [
            {"topic": "Architecture", "color": "blue", "subtopics": ["Hardware", "Software", "Interface"]},
            {"topic": "Components", "color": "green", "subtopics": ["Input", "Processing", "Output"]},
            {"topic": "Testing", "color": "amber", "subtopics": ["Unit Tests", "Integration", "Validation"]},
            {"topic": "Documentation", "color": "purple", "subtopics": ["Specs", "Reports", "Manuals"]},
            {"topic": "Timeline", "color": "coral", "subtopics": ["Phase 1", "Phase 2", "Delivery"]}
        ]
    }

def generate_timeline(project_name, search_results):
    context = '\n\n'.join([f"[{h['filename']}]: {h['content']}" for h in search_results[:6]])
    prompt  = f"""Project: "{project_name}"

Content:
{context[:2500]}

Extract timeline events and milestones from the content.
Return ONLY JSON, no markdown:
{{
  "title": "Project Timeline",
  "events": [
    {{"date": "Jan 2024", "title": "Event title", "description": "Brief description", "type": "milestone"}},
    {{"date": "Feb 2024", "title": "Event title", "description": "Brief description", "type": "task"}}
  ]
}}

Types: milestone, task, issue, completion
Generate 5-8 events. If no specific dates found, use project phases."""
    reply = ask_ai(prompt, 800)
    match = re.search(r'\{[\s\S]*\}', reply)
    if match:
        try:
            return json.loads(match.group())
        except:
            pass
    return {
        "title": f"{project_name} Timeline",
        "events": [
            {"date": "Phase 1", "title": "Project Initiation", "description": "Project started", "type": "milestone"},
            {"date": "Phase 2", "title": "Design", "description": "System design completed", "type": "task"},
            {"date": "Phase 3", "title": "Development", "description": "Implementation phase", "type": "task"},
            {"date": "Phase 4", "title": "Testing", "description": "Testing and validation", "type": "milestone"},
            {"date": "Phase 5", "title": "Delivery", "description": "Project completion", "type": "completion"}
        ]
    }

def generate_executive_briefing(project_name, search_results):
    context = '\n\n'.join([f"[{h['filename']}]: {h['content']}" for h in search_results[:8]])
    prompt  = f"""Project: "{project_name}"

Content:
{context[:3000]}

Generate an executive briefing.
Return ONLY JSON, no markdown:
{{
  "situation": "2-3 sentences describing the situation",
  "key_facts": ["fact1", "fact2", "fact3", "fact4"],
  "risks": ["risk1", "risk2", "risk3"],
  "recommendations": ["recommendation1", "recommendation2", "recommendation3"],
  "action_items": [
    {{"action": "action description", "owner": "Team/Person", "priority": "High/Medium/Low"}}
  ],
  "conclusion": "1-2 sentence conclusion"
}}"""
    reply = ask_ai(prompt, 1000)
    match = re.search(r'\{[\s\S]*\}', reply)
    if match:
        try:
            return json.loads(match.group())
        except:
            pass
    return {
        "situation": f"The {project_name} project requires attention and review.",
        "key_facts": ["Project documentation is available", "Files have been indexed"],
        "risks": ["Timeline adherence", "Resource allocation"],
        "recommendations": ["Review documentation", "Establish clear milestones"],
        "action_items": [{"action": "Review project files", "owner": "Project Lead", "priority": "High"}],
        "conclusion": "Immediate action recommended to ensure project success."
    }

def generate_word_report(project_name, project_desc, search_results):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    context = '\n\n'.join([f"[{h['filename']}]: {h['content']}" for h in search_results[:10]])
    prompt  = f"""Project: "{project_name}"
Description: {project_desc}

Content:
{context[:4000]}

Generate a technical report. Return ONLY JSON, no markdown:
{{
  "executive_summary": "2-3 paragraph executive summary",
  "introduction": "Introduction paragraph",
  "sections": [
    {{"title": "Section Title", "content": "Section content 2-3 paragraphs"}},
    {{"title": "Technical Specifications", "content": "Technical details"}},
    {{"title": "System Architecture", "content": "Architecture description"}},
    {{"title": "Testing and Results", "content": "Test results and findings"}},
    {{"title": "Conclusions", "content": "Conclusions and recommendations"}}
  ]
}}"""

    reply = ask_ai(prompt, 2000)
    match = re.search(r'\{[\s\S]*\}', reply)
    data  = None
    if match:
        try:
            data = json.loads(match.group())
        except:
            pass

    if not data:
        data = {
            "executive_summary": f"This report covers the {project_name} project.",
            "introduction": f"The {project_name} project involves {project_desc or 'technical development'}.",
            "sections": [{"title": "Project Overview", "content": f"Details of {project_name}."}]
        }

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Cover page
    doc.add_paragraph()
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(project_name.upper())
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(10, 22, 40)

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_para.add_run('Technical Project Report')
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(55, 138, 221)

    doc.add_paragraph()
    org_para = doc.add_paragraph()
    org_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    org_para.add_run('Dynalog Defence Organisation').font.size = Pt(12)

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_para.add_run(datetime.date.today().strftime('%B %Y')).font.size = Pt(11)

    doc.add_page_break()

    # Executive summary
    h = doc.add_heading('Executive Summary', level=1)
    h.runs[0].font.color.rgb = RGBColor(10, 22, 40)
    doc.add_paragraph(data.get('executive_summary', ''))

    # Introduction
    h = doc.add_heading('Introduction', level=1)
    h.runs[0].font.color.rgb = RGBColor(10, 22, 40)
    doc.add_paragraph(data.get('introduction', ''))

    # Sections
    for section in data.get('sections', []):
        h = doc.add_heading(section.get('title', 'Section'), level=1)
        h.runs[0].font.color.rgb = RGBColor(10, 22, 40)
        doc.add_paragraph(section.get('content', ''))

    # Footer info
    doc.add_page_break()
    footer_para = doc.add_paragraph()
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run('CONFIDENTIAL — DYNALOG DEFENCE ORGANISATION — INTERNAL USE ONLY')
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(150, 150, 150)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf